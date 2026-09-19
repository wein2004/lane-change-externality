"""Build the project introduction deck as a PowerPoint file.

Content lives in scripts/slides_content.py; this file only renders it.
Each layout is one function so a slide's geometry is readable in one place.

Usage: python scripts/05_build_pptx.py [--out FILE]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from pptx_theme import (  # noqa: E402
    BG, BLUE, CONTENT_W, INK, MARGIN, ORANGE, PANEL, RULE, SLIDE_H, SLIDE_W,
    SUB, TOP, arrow, eyebrow, footnote, hrule, page_number, panel, para_block,
    set_bg, table, textbox, title, write,
)
from slides_content import DECK  # noqa: E402

COLORS = {"INK": INK, "SUB": SUB, "BLUE": BLUE, "ORANGE": ORANGE,
          "PANEL": PANEL, "RULE": RULE}

LEAD_Y = TOP + Inches(1.02)          # clears the title
BODY_TOP = TOP + Inches(1.15)        # slides with no standfirst
BODY_LEAD = TOP + Inches(1.78)       # slides that have one
BODY_H = SLIDE_H - BODY_TOP - Inches(0.95)


def lead(slide, text):
    """The one-line standfirst under a title. Returns the body's top edge."""
    tf = textbox(slide, MARGIN, LEAD_Y, CONTENT_W, Inches(0.6))
    write(tf.paragraphs[0], text, size=13, color=SUB, line=1.45)
    return BODY_LEAD


def card(slide, left, top, width, height, head, body, tone="PANEL", *,
         head_size=14, body_size=12):
    """A labelled block. `tone` picks fill vs. outline treatment."""
    accent = tone in ("BLUE", "ORANGE")
    shape = panel(slide, left, top, width, height,
                  fill=BG if accent else PANEL,
                  line=COLORS[tone] if accent else None)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.22)
    tf.margin_top = tf.margin_bottom = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(tf.paragraphs[0], head, size=head_size, bold=True,
          color=COLORS[tone] if accent else INK, space_after=4)
    if body:
        write(tf.add_paragraph(), body, size=body_size, color=SUB, line=1.35)
    return shape


def fit_image(slide, path, left, top, box_w, box_h):
    """Place an image contained inside a box, centred, on a light panel."""
    with Image.open(ROOT / path) as im:
        iw, ih = im.size
    scale = min(box_w / iw, box_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    panel(slide, left, top, box_w, box_h, fill=PANEL)
    slide.shapes.add_picture(str(ROOT / path),
                             Emu(int(left + (box_w - w) / 2)),
                             Emu(int(top + (box_h - h) / 2)),
                             width=Emu(w), height=Emu(h))


# --------------------------------------------------------------------------- #
# Layouts
# --------------------------------------------------------------------------- #
def L_cover(slide, d):
    eyebrow(slide, d["eyebrow"], top=Inches(2.05))
    tf = textbox(slide, MARGIN, Inches(2.42), CONTENT_W, Inches(1.3))
    write(tf.paragraphs[0], d["title"], size=52, bold=True, line=1.1)
    tf = textbox(slide, MARGIN, Inches(4.0), Inches(8.6), Inches(1.2))
    for i, line in enumerate(d["sub"].split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        write(para, line, size=16, color=SUB, line=1.5)
    hrule(slide, MARGIN, Inches(5.6), Inches(2.4), INK, 2.5)
    footnote(slide, d["foot"])


def L_closing(slide, d):
    tf = textbox(slide, MARGIN, Inches(2.9), CONTENT_W, Inches(1.0))
    write(tf.paragraphs[0], d["title"], size=40, bold=True, line=1.1)
    tf = textbox(slide, MARGIN, Inches(3.9), CONTENT_W, Inches(0.5))
    write(tf.paragraphs[0], d["sub"], size=15, color=SUB)
    hrule(slide, MARGIN, Inches(4.7), Inches(2.4), INK, 2.5)
    footnote(slide, d["foot"])


def L_section(slide, d):
    tf = textbox(slide, MARGIN, Inches(2.5), CONTENT_W, Inches(0.4))
    write(tf.paragraphs[0], d["num"], size=15, color=BLUE, bold=True)
    tf = textbox(slide, MARGIN, Inches(2.95), CONTENT_W, Inches(1.0))
    write(tf.paragraphs[0], d["title"], size=46, bold=True, line=1.1)
    tf = textbox(slide, MARGIN, Inches(4.15), Inches(8.0), Inches(0.6))
    write(tf.paragraphs[0], d["sub"], size=16, color=SUB, line=1.45)


def L_statement(slide, d):
    eyebrow(slide, d["eyebrow"], top=Inches(1.9))
    tf = textbox(slide, MARGIN, Inches(2.3), Inches(10.5), Inches(2.0))
    for i, line in enumerate(d["title"].split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        write(para, line, size=42, bold=True, line=1.18)
    tf = textbox(slide, MARGIN, Inches(4.75), Inches(10.0), Inches(1.0))
    write(tf.paragraphs[0], d["sub"], size=16, color=SUB, line=1.55)


def L_agenda(slide, d):
    title(slide, d["title"])
    rows = d["items"]
    top = BODY_TOP
    h = Inches(0.8)
    for num, name, desc in rows:
        tf = textbox(slide, MARGIN, top + Inches(0.06), Inches(0.7), Inches(0.4))
        write(tf.paragraphs[0], num, size=14, color=BLUE, bold=True)
        tf = textbox(slide, MARGIN + Inches(0.8), top, Inches(2.3), Inches(0.4))
        write(tf.paragraphs[0], name, size=17, bold=True)
        tf = textbox(slide, MARGIN + Inches(3.3), top + Inches(0.05),
                     CONTENT_W - Inches(3.3), Inches(0.45))
        write(tf.paragraphs[0], desc, size=13, color=SUB, line=1.35)
        hrule(slide, MARGIN, top + h - Inches(0.1), CONTENT_W)
        top += h


def L_split(slide, d):
    title(slide, d["title"])
    left_w = Inches(6.6)
    lines = [(t, {"size": s, "color": COLORS[c], "line": 1.55,
                  "space_after": 12}) for t, s, c in d["left"]]
    para_block(slide, MARGIN, BODY_TOP, left_w, lines, height=BODY_H)

    right_x = MARGIN + left_w + Inches(0.6)
    right_w = CONTENT_W - left_w - Inches(0.6)
    cards = d["cards"]
    gap = Inches(0.28)
    h = (Inches(3.7) - gap * (len(cards) - 1)) / len(cards)
    top = BODY_TOP
    for head, body, tone in cards:
        card(slide, right_x, top, right_w, h, head, body, tone)
        top += h + gap
    if d.get("note"):
        footnote(slide, d["note"])


def L_cards3(slide, d):
    title(slide, d["title"])
    top0 = lead(slide, d["lead"]) if d.get("lead") else BODY_TOP
    cards = d["cards"]
    gap = Inches(0.3)
    w = (CONTENT_W - gap * (len(cards) - 1)) / len(cards)
    top = top0 + Inches(0.1)
    for i, (head, sub_line, body) in enumerate(cards):
        left = MARGIN + i * (w + gap)
        shape = panel(slide, left, top, w, Inches(3.35), fill=PANEL)
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.26)
        tf.margin_top = tf.margin_bottom = Inches(0.24)
        write(tf.paragraphs[0], head, size=20, bold=True, space_after=2)
        write(tf.add_paragraph(), sub_line, size=12, color=BLUE, bold=True,
              space_after=10)
        write(tf.add_paragraph(), body, size=13, color=SUB, line=1.45)
    if d.get("note"):
        footnote(slide, d["note"])


def L_metrics(slide, d):
    title(slide, d["title"])
    top0 = lead(slide, d["lead"]) if d.get("lead") else BODY_TOP
    cards = d["cards"]
    gap = Inches(0.3)
    w = (CONTENT_W - gap * (len(cards) - 1)) / len(cards)
    top = top0 + Inches(0.3)
    for i, (value, label, note) in enumerate(cards):
        left = MARGIN + i * (w + gap)
        accent = note.startswith("##")
        shape = panel(slide, left, top, w, Inches(2.75),
                      fill=BG if accent else PANEL,
                      line=ORANGE if accent else None)
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.26)
        tf.margin_top = tf.margin_bottom = Inches(0.26)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        write(tf.paragraphs[0], label, size=13, color=SUB, space_after=8)
        write(tf.add_paragraph(), value, size=34, bold=True, space_after=8)
        write(tf.add_paragraph(), note, size=12, color=SUB)
    if d.get("note"):
        footnote(slide, d["note"])


def L_table(slide, d):
    title(slide, d["title"])
    top = lead(slide, d["lead"]) if d.get("lead") else BODY_TOP
    rows = d["rows"]
    row_h = Inches(0.58) if len(rows) <= 6 else Inches(0.5)
    table(slide, MARGIN, top, CONTENT_W, rows, d["col_w"], row_h=row_h,
          size=12 if len(rows) > 6 else 13)
    if d.get("note"):
        footnote(slide, d["note"])


def L_image(slide, d):
    title(slide, d["title"])
    left_w = Inches(5.3)
    lines = [(t, {"size": s, "color": COLORS[c], "line": 1.55,
                  "space_after": 10}) for t, s, c in d["left"]]
    para_block(slide, MARGIN, BODY_TOP + Inches(0.2), left_w, lines,
               height=Inches(4.0))
    img_x = MARGIN + left_w + Inches(0.5)
    fit_image(slide, d["image"], img_x, BODY_TOP,
              CONTENT_W - left_w - Inches(0.5), Inches(4.3))
    if d.get("note"):
        footnote(slide, d["note"])


def L_image_wide(slide, d):
    title(slide, d["title"])
    fit_image(slide, d["image"], MARGIN, BODY_TOP, CONTENT_W, Inches(3.1))
    lines = [(t, {"size": s, "color": COLORS[c], "line": 1.5,
                  "space_after": 8}) for t, s, c in d["left"]]
    para_block(slide, MARGIN, BODY_TOP + Inches(3.35), CONTENT_W, lines,
               height=Inches(0.9))
    if d.get("note"):
        footnote(slide, d["note"])


def L_pipeline(slide, d):
    title(slide, d["title"])
    steps = d["steps"]
    gap = Inches(0.22)
    arrow_w = Inches(0.3)
    total_gap = (gap * 2 + arrow_w) * (len(steps) - 1)
    w = (CONTENT_W - total_gap) / len(steps)
    top = BODY_TOP + Inches(0.7)
    h = Inches(2.35)
    x = MARGIN
    for i, (num, name, body) in enumerate(steps):
        shape = panel(slide, x, top, w, h, fill=PANEL)
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.18)
        tf.margin_top = tf.margin_bottom = Inches(0.2)
        write(tf.paragraphs[0], num, size=12, color=BLUE, bold=True,
              space_after=6)
        write(tf.add_paragraph(), name, size=17, bold=True, space_after=8)
        for line in body.split("\n"):
            write(tf.add_paragraph(), line, size=11, color=SUB, line=1.3)
        x += w
        if i < len(steps) - 1:
            arrow(slide, x + gap, top + h / 2 - Inches(0.09), arrow_w,
                  Inches(0.18))
            x += gap * 2 + arrow_w
    if d.get("note"):
        footnote(slide, d["note"])


def L_arch(slide, d):
    """The system architecture: four stages, their tools and their outputs."""
    title(slide, d["title"])
    stages = [
        ("資料來源", "NGSIM 公開 API", ["Socrata 端點", "無需驗證", "1,180 萬列"], BLUE),
        ("取得與清理", "01 · src/ngsim", ["鍵集分頁下載", "去重與批次切分", "速度平滑"], None),
        ("事件與對照組", "02_detect_events", ["換道偵測", "四道過濾", "最近鄰配對"], None),
        ("統計與產出", "03 · 04", ["Bootstrap 區間", "OLS 迴歸", "圖表與結果表"], None),
    ]
    gap = Inches(0.26)
    w = (CONTENT_W - gap * (len(stages) - 1)) / len(stages)
    top = BODY_TOP + Inches(0.1)
    h = Inches(3.0)
    for i, (name, tool, items, tone) in enumerate(stages):
        left = MARGIN + i * (w + gap)
        shape = panel(slide, left, top, w, h,
                      fill=BG if tone else PANEL,
                      line=tone)
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.22)
        tf.margin_top = tf.margin_bottom = Inches(0.22)
        write(tf.paragraphs[0], name, size=17, bold=True,
              color=BLUE if tone else INK, space_after=4)
        write(tf.add_paragraph(), tool, size=11, color=SUB, space_after=12)
        for it in items:
            write(tf.add_paragraph(), "·  " + it, size=12, color=INK,
                  line=1.5, space_after=2)
        if i < len(stages) - 1:
            arrow(slide, left + w + Inches(0.04), top + h / 2 - Inches(0.08),
                  Inches(0.18), Inches(0.16))

    # The artefact each stage leaves behind, as a labelled band underneath.
    band_top = top + h + Inches(0.32)
    outs = ["公開資料集", "data/raw  ·  126 MB Parquet",
            "data/processed  ·  事件與對照組", "outputs  ·  圖表、結果表、日誌"]
    for i, text in enumerate(outs):
        left = MARGIN + i * (w + gap)
        tf = textbox(slide, left, band_top, w, Inches(0.4))
        write(tf.paragraphs[0], text, size=11, color=SUB, line=1.35)
    hrule(slide, MARGIN, band_top - Inches(0.14), CONTENT_W)
    footnote(slide, "每個階段都有獨立的輸出與日誌，任何一步都可以單獨重跑，"
                    "不需要重跑整條流程。")


LAYOUTS = {
    "cover": L_cover, "closing": L_closing, "section": L_section,
    "statement": L_statement, "agenda": L_agenda, "split": L_split,
    "cards3": L_cards3, "metrics": L_metrics, "table": L_table,
    "image": L_image, "image_wide": L_image_wide, "pipeline": L_pipeline,
    "arch": L_arch,
}

NO_NUMBER = {"cover", "closing", "section"}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default="變換車道的外部成本.pptx")
    args = ap.parse_args()

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    for i, spec in enumerate(DECK, start=1):
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        LAYOUTS[spec["layout"]](slide, spec)
        if spec["layout"] not in NO_NUMBER:
            page_number(slide, i)

    out = ROOT / args.out
    prs.save(out)
    print(f"wrote {out} ({len(DECK)} slides, {out.stat().st_size / 1e6:.1f} MB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
