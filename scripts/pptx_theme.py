"""Shared drawing helpers for the PowerPoint build.

Keeps `05_build_pptx.py` to layout and content by hiding the python-pptx
boilerplate: CJK-safe fonts, inline emphasis markup, panels, rules and tables.

Inline markup used throughout the deck copy:
    **text**   bold
    @@text@@   bold, accent blue
    ##text##   bold, accent orange
"""

from __future__ import annotations

import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------- #
# Canvas and palette
# --------------------------------------------------------------------------- #
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

MARGIN = Inches(0.85)
TOP = Inches(0.72)
CONTENT_W = SLIDE_W - 2 * MARGIN

BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x11, 0x11)
SUB = RGBColor(0x6B, 0x6B, 0x6B)
RULE = RGBColor(0xE4, 0xE4, 0xE4)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
PANEL = RGBColor(0xF6, 0xF6, 0xF4)

FONT = "Microsoft JhengHei"
FONT_MONO = "Consolas"

MARKUP = re.compile(r"(\*\*.+?\*\*|@@.+?@@|##.+?##)", re.S)


# --------------------------------------------------------------------------- #
# Text
# --------------------------------------------------------------------------- #
def _apply_font(run, name: str) -> None:
    """Set the typeface for latin AND east-asian runs.

    python-pptx only writes the latin typeface. Without the `ea` entry
    PowerPoint falls back to its own default for Chinese glyphs, so the deck
    would render in two different faces.
    """
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def write(
    para,
    text: str,
    *,
    size: int,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = FONT,
    line: float | None = None,
    space_after: int = 0,
) -> None:
    """Write `text` into `para`, honouring the inline emphasis markup."""
    para.space_after = Pt(space_after)
    para.space_before = Pt(0)
    para.alignment = PP_ALIGN.LEFT
    if line is not None:
        para.line_spacing = line

    for chunk in MARKUP.split(text):
        if not chunk:
            continue
        run_bold, run_color, body = bold, color, chunk
        if chunk.startswith("**"):
            run_bold, body = True, chunk[2:-2]
        elif chunk.startswith("@@"):
            run_bold, run_color, body = True, BLUE, chunk[2:-2]
        elif chunk.startswith("##"):
            run_bold, run_color, body = True, ORANGE, chunk[2:-2]
        run = para.add_run()
        run.text = body
        run.font.size = Pt(size)
        run.font.bold = run_bold
        run.font.color.rgb = run_color
        _apply_font(run, font)


def textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para_block(slide, left, top, width, lines, *, height=None, anchor=MSO_ANCHOR.TOP):
    """A stack of paragraphs. `lines` are (text, kwargs) pairs."""
    tf = textbox(slide, left, top, width, height or Inches(1), anchor=anchor)
    for i, (text, kw) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        write(para, text, **kw)
    return tf


# --------------------------------------------------------------------------- #
# Shapes
# --------------------------------------------------------------------------- #
def panel(slide, left, top, width, height, *, fill=PANEL, line=None, line_w=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.text_frame.word_wrap = True
    return shape


def hrule(slide, left, top, width, color=RULE, weight=1.0):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(weight))
    line.shadow.inherit = False
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    return line


def arrow(slide, left, top, width, height, color=RULE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def set_bg(slide, color=BG):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


# --------------------------------------------------------------------------- #
# Composite pieces used on most slides
# --------------------------------------------------------------------------- #
def eyebrow(slide, text, top=None):
    tf = textbox(slide, MARGIN, top or TOP, CONTENT_W, Inches(0.3))
    write(tf.paragraphs[0], text, size=12, color=SUB, bold=True)


def title(slide, text, *, top=None, size=30, color=INK):
    t = top if top is not None else TOP + Inches(0.42)
    tf = textbox(slide, MARGIN, t, CONTENT_W, Inches(0.8))
    write(tf.paragraphs[0], text, size=size, color=color, bold=True, line=1.15)
    return t + Inches(0.42 + 0.22 * (text.count("\n") + 1))


def footnote(slide, text):
    tf = textbox(slide, MARGIN, SLIDE_H - Inches(0.72), CONTENT_W, Inches(0.3))
    write(tf.paragraphs[0], text, size=11, color=SUB)


def page_number(slide, n):
    tf = textbox(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.72),
                 Inches(1), Inches(0.3))
    write(tf.paragraphs[0], str(n), size=11, color=SUB)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT


def table(slide, left, top, width, rows, col_w, *, size=13, header=True,
          row_h=Inches(0.42)):
    """A flat, ruled table. `rows` is a list of lists of cell strings."""
    n_rows, n_cols = len(rows), len(col_w)
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width,
                                   row_h * n_rows)
    tbl = shape.table
    tbl.first_row = header
    tbl.horz_banding = False

    total = sum(col_w)
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Emu(int(width * w / total))

    for r, row in enumerate(rows):
        tbl.rows[r].height = row_h
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if (header and r == 0) else BG
            para = cell.text_frame.paragraphs[0]
            right = text.startswith(">")
            if right:
                text = text[1:]
            write(para, text, size=size,
                  color=INK, bold=(header and r == 0))
            # After write(), which forces LEFT on every paragraph.
            if right:
                para.alignment = PP_ALIGN.RIGHT
    return shape
